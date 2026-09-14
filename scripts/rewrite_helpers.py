# -*- coding: utf-8 -*-
"""就地改写 docx 的保格式工具函数集。

设计原则：复用已有的 run 与 XML 元素，绝不重建段落/单元格，
从而保留源文档的字体、加粗、缩进、表格边框等全部格式。

用法：把这些函数贴进改写脚本，或 from rewrite_helpers import *
"""
import copy
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.table import _Cell


def set_para_text(p, text, bold=None):
    """改写段落文字，保留首个 run 的格式，删除其余 run。

    注意：不要用 p.text = text，那会清空所有 run 级格式。
    """
    runs = p.runs
    if not runs:
        r = p.add_run(text)
        if bold is not None:
            r.font.bold = bold
        return
    r0 = runs[0]
    r0.text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    if bold is not None:
        r0.font.bold = bold


def set_term(p, name, desc):
    """改写「词条名：释义」式段落（run0 加粗词条名，run1 正常释义）。"""
    runs = p.runs
    if len(runs) >= 2:
        runs[0].text = name
        runs[1].text = desc
        for r in runs[2:]:
            r._element.getparent().remove(r._element)
        runs[0].font.bold = True
        runs[1].font.bold = False
    else:
        set_para_text(p, name + desc, bold=False)


def insert_para_after(ref_p, text=None, name=None, desc=None, bold=None):
    """在 ref_p 之后插入同格式段落。返回新 Paragraph。

    坑：新段落完全复制 ref_p 的格式。若 ref_p 已加粗（例如小标题），
    后续段落会误继承加粗——插入正文时要显式传 bold=False。
    """
    new_el = copy.deepcopy(ref_p._element)
    ref_p._element.addnext(new_el)
    np = Paragraph(new_el, ref_p._parent)
    if name is not None:
        set_term(np, name, desc)
    else:
        set_para_text(np, text, bold=bold)
    return np


def _sample_run(tc_element):
    """从表格中取一个「带非空文字」的 run 作为新 run 的格式模板。"""
    tbl = tc_element.getparent().getparent()
    for tr in tbl.findall(qn('w:tr')):
        for tc in tr.findall(qn('w:tc')):
            for pEl in tc.findall(qn('w:p')):
                for rEl in pEl.findall(qn('w:r')):
                    txt = ''.join(t.text or '' for t in rEl.findall(qn('w:t')))
                    if txt.strip():
                        return rEl
    return None


def set_cell_text(cell, text):
    """写单元格文字：只保留首个段落，写入最后一个非空 run 并删除其余 run。

    【重要】本项目的 docx 模板中，单元格常见结构是
        run0 = 空 run（无 rPr，字体为空） + run1 = 带新宋体/10.5 的真实文字
    若写入 run0 并删除其余 run，会整体丢失字体和字号。
    因此必须定位到「最后一个文字非空的 run」再写入。
    """
    ps = cell.paragraphs
    if not ps:
        return
    p0 = ps[0]
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)

    if not p0.runs:
        sample = _sample_run(cell._tc)
        if sample is None:
            p0.add_run(text)
            return
        new_r = copy.deepcopy(sample)
        for t in new_r.findall(qn('w:t')):
            new_r.remove(t)
        p0._p.append(new_r)

    runs = p0.runs
    target = None
    for r in runs:
        if r.text.strip():
            target = r
    if target is None:
        target = runs[-1]
    for r in runs:
        if r is not target:
            r._element.getparent().remove(r._element)
    target.text = text


def set_row(table, ri, values):
    """按列序重写整行。加列/改列序后必须调用它，否则列序会错。"""
    for ci, v in enumerate(values):
        set_cell_text(table.rows[ri].cells[ci], v)


def delete_row(table, idx):
    tr = table.rows[idx]._tr
    tr.getparent().remove(tr)


def insert_row_after(table, ref_idx, values=None):
    """在 ref_idx 行之后插入同格式新行并填入 values。"""
    ref_tr = table.rows[ref_idx]._tr
    new_tr = copy.deepcopy(ref_tr)
    ref_tr.addnext(new_tr)
    if values:
        for ci, v in enumerate(values):
            set_cell_text(table.rows[ref_idx + 1].cells[ci], v)
    return table.rows[ref_idx + 1]


def add_column(table, values):
    """在表格末尾追加一列，values 为逐行文本（含表头）。

    注意：新列追加在最后，调用后必须用 set_row 按目标列序重写每一行，
    否则列序会是 [原列..., 新列] 而不是期望的位置。
    追加的单元格复制同行最后一格的格式，表头行的加粗会自动保持。
    """
    tbl = table._tbl
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        gcs = grid.findall(qn('w:gridCol'))
        new_gc = copy.deepcopy(gcs[-1])
        gcs[-1].addnext(new_gc)
    for ti, tr in enumerate(tbl.findall(qn('w:tr'))):
        tcs = tr.findall(qn('w:tc'))
        new_tc = copy.deepcopy(tcs[-1])
        tcs[-1].addnext(new_tc)
        set_cell_text(_Cell(new_tc, table), values[ti] if ti < len(values) else '')


def force_cell_font(table, name, size_pt, bold_map=None):
    """兜底修复表格字体：bold_map = {行号: 是否加粗}，缺省行保持原值。

    用于历史文档已被错误改写、字体丢失后的补救。
    """
    for ri, row in enumerate(table.rows):
        for c in row.cells:
            for p in c.paragraphs:
                vis = [r for r in p.runs if r.text.strip()]
                r = vis[-1] if vis else (p.runs[-1] if p.runs else None)
                if r is None:
                    continue
                r.font.name = name
                if r._element.rPr is not None and r._element.rPr.rFonts is not None:
                    r._element.rPr.rFonts.set(qn('w:eastAsia'), name)
                r.font.size = Pt(size_pt)
                if bold_map and ri in bold_map:
                    r.font.bold = bold_map[ri]


def verify_cell_fonts(path, name='新宋体', size_pt=10.5):
    """校验表格字体是否统一（默认新宋体/10.5），返回异常单元格列表。

    注意：不同文档基准不同——概要/详细设计用新宋体，接口/数据库/运维/原型用宋体；
    封面「修改记录表」常用更大字号（11/12 号），属正常差异，需人工排除。
    先用 inspect 确认本文档的实际字体基准，再传入对应参数。
    """
    from docx import Document
    doc = Document(path)
    bad = []
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            for ci, c in enumerate(row.cells):
                for p in c.paragraphs:
                    vis = [r for r in p.runs if r.text.strip()]
                    if not vis:
                        continue
                    r = vis[0]
                    f = r.font
                    if f.name != name or (f.size and f.size.pt != size_pt):
                        bad.append((ti, ri, ci, f.name, f.size.pt if f.size else None))
    return bad


# ---------------------------------------------------------------------------
# 对齐模板标题体系：多级自动编号注入 + 模板风格标题
# ---------------------------------------------------------------------------

def add_numbering_part(doc, template_docx):
    """把模板的 word/numbering.xml（多级列表定义）注入目标文档。

    目标文档常没有 numbering 部件，此时给段落设 w:numPr 不会渲染出编号。
    注入后所有标题应使用同一个 numId，否则同级计数会重启。
    """
    import zipfile
    from docx.oxml import parse_xml
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT, CONTENT_TYPE as CT

    blob = zipfile.ZipFile(template_docx).read("word/numbering.xml")
    try:
        from docx.parts.numbering import NumberingPart
        part = NumberingPart(PackURI("/word/numbering.xml"), CT.WML_NUMBERING,
                             parse_xml(blob), doc.part.package)
    except Exception:
        part = Part(PackURI("/word/numbering.xml"), CT.WML_NUMBERING, blob,
                    doc.part.package)
    doc.part.relate_to(part, RT.NUMBERING)
    return part


def apply_template_heading(p, level, text, num_id="2",
                           l1_font="黑体", l2_font="仿宋", size_pt=16.0):
    """把段落改造成模板风格标题：多级自动编号 + outlineLvl + 指定字体字号。

    level: 0=大标题（黑体），1=小标题（仿宋）。标题文字不再包含手动编号。
    注意：本机常未安装 仿宋_GB2312，用 仿宋 替代。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    runs = p.runs
    if not runs:
        p.add_run(text)
    else:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)

    font = l1_font if level == 0 else l2_font
    for r in p.runs:
        r.font.name = font
        r.font.size = Pt(size_pt)
        r.font.bold = False
        rPr = r._element.get_or_add_rPr()
        rf = rPr.get_or_add_rFonts()
        for a in ("w:ascii", "w:eastAsia", "w:hAnsi", "w:cs"):
            rf.set(qn(a), font)
        rf.set(qn("w:hint"), "eastAsia")

    pPr = p._p.get_or_add_pPr()
    for tag in ("w:pStyle", "w:numPr", "w:outlineLvl", "w:jc", "w:ind",
                "w:spacing", "w:rPr"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)

    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), str(level))
    numId = OxmlElement("w:numId"); numId.set(qn("w:val"), str(num_id))
    numPr.append(ilvl); numPr.append(numId)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "left")
    ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), str(level))
    prPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:eastAsia", "w:hAnsi", "w:cs"):
        rf.set(qn(a), font)
    rf.set(qn("w:hint"), "eastAsia")
    bCs = OxmlElement("w:bCs")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size_pt * 2)))
    prPr.append(rf); prPr.append(bCs); prPr.append(sz)
    for e in (numPr, jc, ol, prPr):     # 必须按 schema 顺序追加
        pPr.append(e)


def place_before_sectPr(body, el):
    """把元素放到正文末尾（sectPr 之前）。直接 body.append 会破坏文档结构。"""
    from docx.oxml.ns import qn
    sectPr = body.find(qn("w:sectPr"))
    if sectPr is not None:
        sectPr.addprevious(el)
    else:
        body.append(el)


def append_multilevel_definition(doc, num_id=None, abs_id=None):
    """向文档**已有**的 numbering 部件追加一套多级编号定义（%1. / %1.%2. / ...）。

    关键坑：目标文档可能已有 numbering 部件，且常见 numId（如 2）可能指向
    空 lvlText 的项目符号列表——直接用会渲染不出编号。必须先列出现有
    numId / abstractNumId，用未占用的 ID 追加新定义，再把标题指过去。
    返回 (num_id, abs_id)。
    """
    import re
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    npart_el = doc.part.numbering_part.element
    xml = re.sub(r' xmlns:[a-zA-Z0-9]+="[^"]*"', '',
                 npart_el.xml) if False else None
    used_nums = [int(e.get(qn('w:numId'))) for e in npart_el.findall(qn('w:num'))]
    used_abs = [int(e.get(qn('w:abstractNumId')))
                for e in npart_el.findall(qn('w:abstractNum'))]
    num_id = num_id or max(used_nums or [0]) + 10
    abs_id = abs_id or max(used_abs or [0]) + 10

    lvls = []
    for i in range(9):
        text = ".".join("%%%d" % (j + 1) for j in range(i + 1)) + "."
        ind = 360 * i
        lvls.append(
            f'<w:lvl w:ilvl="{i}">'
            f'<w:start w:val="1"/>'
            f'<w:numFmt w:val="decimal"/>'
            f'<w:lvlText w:val="{text}"/>'
            f'<w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{ind}" w:hanging="{ind}"/></w:pPr>'
            f'</w:lvl>')
    abs_el = parse_xml(
        f'<w:abstractNum xmlns:w="{W}" w:abstractNumId="{abs_id}">'
        '<w:multiLevelType w:val="hybridMultilevel"/>' + "".join(lvls) +
        '</w:abstractNum>')
    num_el = parse_xml(
        f'<w:num xmlns:w="{W}" w:numId="{num_id}">'
        f'<w:abstractNumId w:val="{abs_id}"/></w:num>')
    first_num = npart_el.find(qn('w:num'))
    if first_num is not None:
        first_num.addprevious(abs_el)   # schema 顺序：abstractNum 全部在 num 之前
    else:
        npart_el.append(abs_el)
    npart_el.append(num_el)
    return str(num_id), str(abs_id)
